"""Builds the Published_Docs .docx deliverables from their User_Docs Markdown sources.

Run at a release (CLAUDE.md: Published_Docs are refreshed then, not between releases):

    python Published_Docs/build_published_docs.py

Needs python-docx (pip install python-docx). Handles the Markdown these guides use:
headings, paragraphs, **bold**, `code`, [links](...), bullet and numbered lists (nested
by indentation), fenced code blocks and pipe tables. Anything else is written as text.
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

DOCS = [
    # (Markdown source, .docx output, title, subtitle)
    ('User_Docs/Admin_Installation.md', 'Published_Docs/Installation_Guide.docx',
     'Installation Guide', 'BlueTrack: CyberArk PAM Blueprint progress tracking'),
    ('User_Docs/Developer_HostSetup.md', 'Published_Docs/Development_Host_Setup_Guide.docx',
     'Development Host Setup Guide', 'BlueTrack: CyberArk PAM Blueprint progress tracking'),
    ('User_Docs/User_GettingStarted.md', 'Published_Docs/Getting_Started_Guide.docx',
     'Getting Started with BlueTrack', 'BlueTrack: CyberArk PAM Blueprint progress tracking'),
    ('User_Docs/Admin_SettingsReference.md', 'Published_Docs/Administrator_Settings_Guide.docx',
     'Administrator Settings Guide', 'BlueTrack: CyberArk PAM Blueprint progress tracking'),
    ('User_Docs/Release_Notes.md', 'Published_Docs/Release_Notes.docx',
     'Release Notes', 'BlueTrack: CyberArk PAM Blueprint progress tracking'),
]

INLINE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')


def add_inline(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
        elif part.startswith('['):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', part)
            paragraph.add_run(m.group(1))
        else:
            paragraph.add_run(part)


def list_style(document, numbered, level):
    name = ('List Number' if numbered else 'List Bullet') + ('' if level == 0 else f' {min(level + 1, 3)}')
    return document.styles[name]


def build(source, output, title, subtitle):
    lines = open(os.path.join(ROOT, source), encoding='utf-8').read().splitlines()
    document = Document()
    document.styles['Normal'].font.name = 'Calibri'
    document.styles['Normal'].font.size = Pt(11)

    document.add_paragraph(title, style='Title')
    sub = document.add_paragraph(subtitle)
    sub.runs[0].italic = True
    sub.alignment = WD_ALIGN_PARAGRAPH.LEFT

    i = 0
    skipped_h1 = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                p = document.add_paragraph()
                run = p.add_run(lines[i].rstrip() or ' ')
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Pt(18)
                i += 1
            i += 1
            continue

        if stripped.startswith('|') and i + 1 < len(lines) and re.match(r'^\s*\|[\s:|-]+\|\s*$', lines[i + 1]):
            header = [c.strip() for c in stripped.strip('|').split('|')]
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            table = document.add_table(rows=1, cols=len(header))
            table.style = 'Table Grid'
            for cell, text in zip(table.rows[0].cells, header):
                cell.paragraphs[0].add_run(text.replace('**', '')).bold = True
            for row in rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    add_inline(cell.paragraphs[0], text)
            document.add_paragraph()
            continue

        heading = re.match(r'^(#{1,4})\s+(.*)$', stripped)
        if heading:
            level = len(heading.group(1))
            if level == 1 and not skipped_h1:
                skipped_h1 = True  # the title above replaces the Markdown's own title
            else:
                document.add_heading(heading.group(2).replace('**', ''), level=max(1, level - 1))
            i += 1
            continue

        item = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)$', line)
        if item:
            level = len(item.group(1).replace('\t', '    ')) // 2
            p = document.add_paragraph(style=list_style(document, item.group(2)[0].isdigit(), min(level, 2)))
            text = item.group(3)
            # A list item's text can continue on indented lines.
            while i + 1 < len(lines) and lines[i + 1].startswith('   ') and lines[i + 1].strip() \
                    and not re.match(r'^\s*([-*]|\d+\.)\s+', lines[i + 1]) and not lines[i + 1].strip().startswith('```'):
                i += 1
                text += ' ' + lines[i].strip()
            add_inline(p, text)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        text = stripped
        while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r'^(\s*([-*]|\d+\.)\s+|#{1,4}\s|```|\|)', lines[i + 1].strip()):
            i += 1
            text += ' ' + lines[i].strip()
        if text.startswith('>'):
            text = text.lstrip('> ')
        add_inline(document.add_paragraph(), text)
        i += 1

    document.save(os.path.join(ROOT, output))
    print(f'{source} -> {output}')


if __name__ == '__main__':
    for doc in DOCS:
        build(*doc)
    sys.exit(0)
